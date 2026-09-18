import os
import re
import html
import urllib.request
from typing import List, Dict, Any, Tuple

def clean_text(text: str) -> str:
    if not text:
        return ""
    # Normalize whitespaces and line breaks
    text = re.sub(r'[\r\t]', ' ', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    text = re.sub(r' +', ' ', text)
    return text.strip()

def extract_text_from_url(url: str) -> Tuple[str, List[Dict[str, Any]]]:
    try:
        req = urllib.request.Request(
            url,
            headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'}
        )
        with urllib.request.urlopen(req, timeout=12) as response:
            html_content = response.read().decode('utf-8', errors='ignore')

        # Remove script, style, svg, and noscript elements
        clean_html = re.sub(r'<(script|style|svg|noscript)[^>]*>.*?</\1>', '', html_content, flags=re.DOTALL | re.IGNORECASE)
        # Convert paragraph/break tags to newlines
        clean_html = re.sub(r'<(p|br|div|h1|h2|h3|h4|li)[^>]*>', '\n', clean_html, flags=re.IGNORECASE)
        # Strip all HTML tags
        text = re.sub(r'<[^>]+>', ' ', clean_html)
        # Unescape HTML entities
        text = html.unescape(text)
        cleaned = clean_text(text)
        
        if not cleaned:
            cleaned = f"URL content fetched from {url}, but no plain text was found."
            
        pages = [{"page_number": 1, "content": cleaned[:2000]}]
        return cleaned, pages
    except Exception as e:
        err_msg = f"Failed to fetch or process content from URL ({url}): {str(e)}"
        return err_msg, [{"page_number": 1, "content": err_msg}]

def extract_text_from_file(file_path: str, filename: str) -> Tuple[str, List[Dict[str, Any]]]:
    ext = os.path.splitext(filename)[1].lower()
    full_text = ""
    pages_or_slides: List[Dict[str, Any]] = []
    
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            for idx, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                cleaned = clean_text(page_text)
                if cleaned:
                    pages_or_slides.append({
                        "page_number": idx + 1,
                        "content": cleaned
                    })
                    full_text += f"\n--- Page {idx + 1} ---\n" + cleaned
                    
        elif ext in [".docx", ".doc"]:
            import docx
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            full_text = "\n\n".join(paragraphs)
            pages_or_slides.append({
                "page_number": 1,
                "content": full_text
            })
            
        elif ext in [".pptx", ".ppt"]:
            from pptx import Presentation
            prs = Presentation(file_path)
            for idx, slide in enumerate(prs.slides):
                slide_texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        slide_texts.append(shape.text)
                slide_content = clean_text("\n".join(slide_texts))
                if slide_content:
                    pages_or_slides.append({
                        "page_number": idx + 1,
                        "content": slide_content
                    })
                    full_text += f"\n--- Slide {idx + 1} ---\n" + slide_content
                    
        else: # txt, md, json, csv
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                full_text = f.read()
            pages_or_slides.append({
                "page_number": 1,
                "content": full_text
            })
    except Exception as e:
        full_text = f"Error extracting document contents: {str(e)}"
        pages_or_slides.append({"page_number": 1, "content": full_text})
        
    return clean_text(full_text), pages_or_slides

def chunk_text(text: str, chunk_size: int = 800, overlap: int = 150) -> List[str]:
    cleaned = clean_text(text)
    if not cleaned:
        return []
    
    chunks = []
    start = 0
    while start < len(cleaned):
        end = min(start + chunk_size, len(cleaned))
        chunk = cleaned[start:end]
        chunks.append(chunk)
        if end == len(cleaned):
            break
        start += chunk_size - overlap
    return chunks

def search_relevant_chunks(chunks: List[str], query: str, top_k: int = 3) -> List[str]:
    if not chunks:
        return []
    if not query:
        return chunks[:top_k]
        
    query_words = set(re.findall(r'\w+', query.lower()))
    scored_chunks = []
    for chunk in chunks:
        chunk_words = set(re.findall(r'\w+', chunk.lower()))
        overlap_score = len(query_words.intersection(chunk_words))
        scored_chunks.append((overlap_score, chunk))
        
    scored_chunks.sort(key=lambda x: x[0], reverse=True)
    return [c[1] for c in scored_chunks[:top_k]]

